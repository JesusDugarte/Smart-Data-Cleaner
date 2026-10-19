
import os
import subprocess
import sys
import tempfile
import unittest
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook

RAIZ_PROYECTO = Path(__file__).resolve().parent.parent


class PruebaFidelidadDatosExportados(unittest.TestCase):

    def test_fidelidad_csv_y_excel(self):

        registros_esperados = [
            (30, "Producto AA", 100.0, "archivo_integracion_L.csv"),
            (31, "Producto AB", 200.0, "archivo_integracion_L.csv"),
            (32, "Producto AC", 300.0, "archivo_integracion_L.csv"),
            (33, "Producto AD", 400.0, "archivo_integracion_L2.csv"),
            (34, "Producto AE", 500.0, "archivo_integracion_L2.csv"),
            (35, "Producto AF", 600.0, "archivo_integracion_L2.csv"),
        ]

        archivos_entrada = [
            RAIZ_PROYECTO / "pruebas_v2_06" / "archivo_integracion_L.csv",
            RAIZ_PROYECTO / "pruebas_v2_06" / "archivo_integracion_L2.csv",
        ]

        archivo_principal = RAIZ_PROYECTO / "inicio_v2_06.py"

        for archivo in archivos_entrada:
            self.assertTrue(
                archivo.is_file(),
                f"No existe el archivo de entrada: {archivo.name}",
            )

        self.assertTrue(archivo_principal.is_file())

        with tempfile.TemporaryDirectory(
            prefix="prueba_fidelidad_"
        ) as carpeta_temporal:

            carpeta = Path(carpeta_temporal)
            ruta_salida = carpeta / "salida"
            ruta_salida.mkdir()

            contenido = archivo_principal.read_text(
                encoding="utf-8-sig"
            )

            bloque_original = (
                'archivos_csv = [\n'
                '    Path("pruebas_v2_06/archivo_X.csv"),\n'
                '    Path("pruebas_v2_06/archivo_W.csv")\n'
                ']'
            )

            bloque_prueba = (
                'archivos_csv = [\n'
                f'    Path("{archivos_entrada[0].as_posix()}"),\n'
                f'    Path("{archivos_entrada[1].as_posix()}")\n'
                ']'
            )

            self.assertIn(
                bloque_original,
                contenido,
                "No se encontró la configuración de entrada esperada.",
            )

            contenido = contenido.replace(
                bloque_original,
                bloque_prueba,
                1,
            )

            rutas_salida = {
                '"salida/datos_limpios.csv"':
                    ruta_salida / "datos_limpios.csv",
                '"salida/informe_proceso.txt"':
                    ruta_salida / "informe_proceso.txt",
                '"salida/informe_proceso.xlsx"':
                    ruta_salida / "informe_proceso.xlsx",
            }

            for ruta_original, ruta_temporal in rutas_salida.items():
                contenido = contenido.replace(
                    ruta_original,
                    f'"{ruta_temporal.as_posix()}"',
                )

            script_temporal = carpeta / "inicio_prueba.py"
            script_temporal.write_text(
                contenido,
                encoding="utf-8",
            )

            entorno = os.environ.copy()
            entorno["PYTHONPATH"] = str(RAIZ_PROYECTO)
            entorno["PYTHONIOENCODING"] = "utf-8"

            resultado = subprocess.run(
                [sys.executable, str(script_temporal)],
                cwd=RAIZ_PROYECTO,
                capture_output=True,
                text=True,
                encoding="utf-8",
                errors="replace",
                env=entorno,
                timeout=60,
            )

            salida_consola = resultado.stdout + "\n" + resultado.stderr

            self.assertEqual(
                resultado.returncode,
                0,
                "El proceso terminó con error:\n" + salida_consola,
            )

            archivo_csv = ruta_salida / "datos_limpios.csv"
            informe_txt = ruta_salida / "informe_proceso.txt"
            informe_excel = ruta_salida / "informe_proceso.xlsx"

            for archivo in (archivo_csv, informe_txt, informe_excel):
                self.assertTrue(
                    archivo.is_file(),
                    f"No se generó el archivo esperado: {archivo.name}",
                )

            datos_csv = pd.read_csv(archivo_csv)

            columnas_esperadas = [
                "ID",
                "Nombre",
                "Precio",
                "Archivo_Origen",
            ]

            self.assertEqual(
                datos_csv.columns.tolist(),
                columnas_esperadas,
                "Las columnas del CSV no son las esperadas.",
            )

            registros_csv = list(
                datos_csv[
                    columnas_esperadas
                ].itertuples(index=False, name=None)
            )

            self.assertEqual(
                registros_csv,
                registros_esperados,
                "Los datos del CSV no coinciden con los registros esperados.",
            )

            libro = load_workbook(
                informe_excel,
                read_only=True,
                data_only=True,
            )

            try:
                self.assertIn(
                    "Datos_Limpios",
                    libro.sheetnames,
                    "No existe la hoja Datos_Limpios en el Excel.",
                )

                hoja = libro["Datos_Limpios"]

                columnas_excel = [
                    celda.value
                    for celda in hoja[3]
                ]

                columnas_excel = columnas_excel[
                    :len(columnas_esperadas)
                ]

                self.assertEqual(
                    columnas_excel,
                    columnas_esperadas,
                    "Las columnas de Datos_Limpios no son las esperadas.",
                )

                registros_excel = [
                    tuple(fila)
                    for fila in hoja.iter_rows(
                        min_row=4,
                        max_col=len(columnas_esperadas),
                        values_only=True,
                    )
                ]

                self.assertEqual(
                    registros_excel,
                    registros_esperados,
                    "Los datos del Excel no coinciden con los registros esperados.",
                )

            finally:
                libro.close()

if __name__ == "__main__":
    unittest.main(verbosity=2)
