import os
import subprocess
import sys
import tempfile
import unittest
from pathlib import Path

from openpyxl import load_workbook


RAIZ_PROYECTO = Path(__file__).resolve().parent.parent


class PruebaBloqueoIntegral(unittest.TestCase):

    def test_bloqueo_no_exporta_datos_y_genera_informes(self):
        archivo_x = RAIZ_PROYECTO / "pruebas_v2_06" / "archivo_X.csv"
        archivo_z = RAIZ_PROYECTO / "pruebas_v2_06" / "archivo_Z.csv"
        archivo_principal = RAIZ_PROYECTO / "inicio_v2_06.py"

        self.assertTrue(archivo_x.is_file())
        self.assertTrue(archivo_z.is_file())
        self.assertTrue(archivo_principal.is_file())

        with tempfile.TemporaryDirectory(
            prefix="prueba_bloqueo_integral_"
        ) as carpeta_temporal:

            carpeta = Path(carpeta_temporal)
            ruta_salida = str(carpeta).replace("\\", "/")

            contenido = archivo_principal.read_text(encoding="utf-8-sig")

            bloque_original = (
                'archivos_csv = [\n'
                '    Path("pruebas_v2_06/archivo_X.csv"),\n'
                '    Path("pruebas_v2_06/archivo_W.csv")\n'
                ']'
            )

            bloque_prueba = (
                'archivos_csv = [\n'
                f'    Path("{archivo_x.as_posix()}"),\n'
                f'    Path("{archivo_z.as_posix()}")\n'
                ']'
            )

            self.assertIn(
                bloque_original,
                contenido,
                "No se encontró la configuración de entrada esperada."
            )

            contenido = contenido.replace(
                bloque_original,
                bloque_prueba,
                1
            )

            contenido = contenido.replace(
                '"salida/datos_limpios.csv"',
                f'"{ruta_salida}/datos_limpios.csv"'
            )
            contenido = contenido.replace(
                '"salida/informe_proceso.txt"',
                f'"{ruta_salida}/informe_proceso.txt"'
            )
            contenido = contenido.replace(
                '"salida/informe_proceso.xlsx"',
                f'"{ruta_salida}/informe_proceso.xlsx"'
            )

            script_temporal = carpeta / "inicio_prueba.py"
            script_temporal.write_text(
                contenido,
                encoding="utf-8"
            )

            ejecucion = subprocess.run(
                [sys.executable, str(script_temporal)],
                cwd=RAIZ_PROYECTO,
                capture_output=True,
                text=True,
                encoding="utf-8",
                errors="replace",
                env={
                    **os.environ,
                    "PYTHONPATH": (
                        str(RAIZ_PROYECTO)
                        + os.pathsep
                        + os.environ.get("PYTHONPATH", "")
                    ),
                    "PYTHONIOENCODING": "utf-8"
                }
            )

            salida_consola = ejecucion.stdout + ejecucion.stderr

            self.assertEqual(
                ejecucion.returncode,
                0,
                f"El programa terminó con error:\n{salida_consola}"
            )

            self.assertIn("Estado final: BLOQUEADO", salida_consola)
            self.assertIn("Consolidación realizada: False", salida_consola)

            self.assertFalse(
                (carpeta / "datos_limpios.csv").exists(),
                "Se generó un CSV limpio a pesar del bloqueo."
            )

            archivo_txt = carpeta / "informe_proceso.txt"
            archivo_excel = carpeta / "informe_proceso.xlsx"

            self.assertTrue(archivo_txt.is_file())
            self.assertTrue(archivo_excel.is_file())

            texto_informe = archivo_txt.read_text(encoding="utf-8")
            self.assertIn("BLOQUEADO", texto_informe)
            self.assertIn("NO_REALIZADA", texto_informe)
            self.assertIn(
                "MOTIVO DEL BLOQUEO",
                texto_informe
            )

            self.assertIn(
                "Se detectaron incompatibilidades de tipos.",
                texto_informe
            )

            self.assertIn(
                "Columna: Precio",
                texto_informe
            )

            self.assertIn(
                "ETAPAS NO EJECUTADAS",
                texto_informe
            )


            libro = load_workbook(
                archivo_excel,
                read_only=True,
                data_only=True
            )

            try:
                self.assertIn("Resumen", libro.sheetnames)

                valores_excel = [
                    str(celda.value)
                    for hoja in libro.worksheets
                    for fila in hoja.iter_rows()
                    for celda in fila
                    if celda.value is not None
                ]

                self.assertIn("BLOQUEADO", valores_excel)
                self.assertIn("NO_REALIZADA", valores_excel)
            finally:
                libro.close()


if __name__ == "__main__":
    unittest.main(verbosity=2)
