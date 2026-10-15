import os
import subprocess
import sys
import tempfile
import unittest
from pathlib import Path

from openpyxl import load_workbook


RAIZ_PROYECTO = Path(__file__).resolve().parent.parent


class PruebaIntegracionEncabezadosDuplicados(unittest.TestCase):

    def test_bloqueo_por_encabezados_duplicados(self):

        archivo_duplicados = (
            RAIZ_PROYECTO
            / "pruebas_v2_06"
            / "prueba_encabezados_duplicados.csv"
        )

        archivo_valido = (
            RAIZ_PROYECTO
            / "pruebas_v2_06"
            / "archivo_integracion_L.csv"
        )

        archivo_principal = RAIZ_PROYECTO / "inicio_v2_06.py"

        self.assertTrue(
            archivo_duplicados.is_file(),
            "No existe el CSV con encabezados duplicados."
        )

        self.assertTrue(
            archivo_valido.is_file(),
            "No existe el CSV válido de integración."
        )

        self.assertTrue(
            archivo_principal.is_file(),
            "No existe inicio_v2_06.py."
        )

        with tempfile.TemporaryDirectory(
            prefix="prueba_encabezados_duplicados_"
        ) as carpeta_temporal:

            carpeta = Path(carpeta_temporal)
            ruta_salida = carpeta / "salida"
            ruta_salida.mkdir()

            contenido = archivo_principal.read_text(
                encoding="utf-8-sig"
            )

            bloque_original = (
                'archivos_csv = [\n'
                '    Path("pruebas_v2_06/archivo_X.csv"),\n'
                '    Path("pruebas_v2_06/archivo_W.csv")\n'
                ']'
            )

            bloque_prueba = (
                'archivos_csv = [\n'
                f'    Path("{archivo_duplicados.as_posix()}"),\n'
                f'    Path("{archivo_valido.as_posix()}")\n'
                ']'
            )

            self.assertIn(
                bloque_original,
                contenido,
                "No se encontró la configuración de entrada esperada."
            )

            contenido = contenido.replace(
                bloque_original,
                bloque_prueba,
                1
            )

            contenido = contenido.replace(
                '"salida/datos_limpios.csv"',
                f'"{(ruta_salida / "datos_limpios.csv").as_posix()}"'
            )

            contenido = contenido.replace(
                '"salida/informe_proceso.txt"',
                f'"{(ruta_salida / "informe_proceso.txt").as_posix()}"'
            )

            contenido = contenido.replace(
                '"salida/informe_proceso.xlsx"',
                f'"{(ruta_salida / "informe_proceso.xlsx").as_posix()}"'
            )

            script_temporal = carpeta / "inicio_prueba.py"

            script_temporal.write_text(
                contenido,
                encoding="utf-8"
            )

            entorno = os.environ.copy()
            entorno["PYTHONPATH"] = str(RAIZ_PROYECTO)
            entorno["PYTHONIOENCODING"] = "utf-8"

            resultado = subprocess.run(
                [sys.executable, str(script_temporal)],
                cwd=RAIZ_PROYECTO,
                capture_output=True,
                text=True,
                encoding="utf-8",
                errors="replace",
                env=entorno,
                timeout=60
            )

            salida_consola = (
                resultado.stdout + "\n" + resultado.stderr
            )

            self.assertEqual(
                resultado.returncode,
                0,
                f"El programa terminó con error:\n{salida_consola}"
            )

            self.assertIn(
                "Estado de consolidación: BLOQUEADO",
                salida_consola
            )

            self.assertIn(
                "No se realizará la consolidación porque hay encabezados duplicados.",
                salida_consola
            )

            self.assertIn(
                "Archivo afectado: prueba_encabezados_duplicados.csv",
                salida_consola
            )

            self.assertIn(
                "Columnas duplicadas:",
                salida_consola
            )

            archivo_limpio = ruta_salida / "datos_limpios.csv"
            informe_txt = ruta_salida / "informe_proceso.txt"
            informe_excel = ruta_salida / "informe_proceso.xlsx"

            self.assertFalse(
                archivo_limpio.exists(),
                "No debe exportarse un CSV limpio cuando la consolidación está bloqueada."
            )

            self.assertTrue(
                informe_txt.is_file(),
                "No se generó el informe de texto."
            )

            self.assertTrue(
                informe_excel.is_file(),
                "No se generó el informe de Excel."
            )

            contenido_txt = informe_txt.read_text(
                encoding="utf-8"
            )

            self.assertIn(
                "BLOQUEADO",
                contenido_txt
            )

            self.assertIn(
                "NO_REALIZADA",
                contenido_txt
            )

            self.assertIn(
                "prueba_encabezados_duplicados.csv",
                contenido_txt
            )

            self.assertIn(
                "encabezados duplicados",
                contenido_txt.lower()
            )

            libro = load_workbook(
                informe_excel,
                read_only=True,
                data_only=True
            )

            try:
                self.assertIn(
                    "Resumen",
                    libro.sheetnames,
                    "El informe Excel no contiene la hoja Resumen."
                )

                contenido_excel = "\n".join(
                    str(valor)
                    for fila in libro["Resumen"].iter_rows(
                        values_only=True
                    )
                    for valor in fila
                    if valor is not None
                )

                self.assertIn(
                    "BLOQUEADO",
                    contenido_excel
                )

                self.assertIn(
                    "NO_REALIZADA",
                    contenido_excel
                )

            finally:
                libro.close()

if __name__ == "__main__":
    unittest.main()
