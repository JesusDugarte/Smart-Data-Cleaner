
import os
import subprocess
import sys
import tempfile
import unittest
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook


RAIZ_PROYECTO = Path(__file__).resolve().parent.parent


class PruebaIntegracionCalidad(unittest.TestCase):

    def test_correcciones_calidad_de_principio_a_fin(self):

        with tempfile.TemporaryDirectory(
            prefix="prueba_integracion_calidad_"
        ) as carpeta_temporal:

            carpeta = Path(carpeta_temporal)

            archivo_A = carpeta / "prueba_A.csv"
            archivo_B = carpeta / "prueba_B.csv"

            archivo_A.write_text(
                "ID,Nombre,Precio\n"
                '901,"  Producto ZZ  ",100.0\n'
                "902,Producto ZZ,200.0\n"
                "903,PRODUCTO ZZ,300.0\n",
                encoding="utf-8"
            )

            archivo_B.write_text(
                "ID,Nombre,Precio\n"
                "904,Producto ZZ,400.0\n"
                "905,Producto YY,500.0\n",
                encoding="utf-8"
            )

            self.assertTrue(archivo_A.is_file())
            self.assertTrue(archivo_B.is_file())

            datos_A = pd.read_csv(archivo_A)
            datos_B = pd.read_csv(archivo_B)

            self.assertEqual(len(datos_A), 3)
            self.assertEqual(len(datos_B), 2)
            self.assertEqual(
                len(datos_A) + len(datos_B),
                5
            )

            archivo_principal = (
                RAIZ_PROYECTO / "inicio_v2_06.py"
            )

            self.assertTrue(
                archivo_principal.is_file(),
                "No existe inicio_v2_06.py."
            )

            ruta_salida = carpeta / "salida"
            ruta_salida.mkdir()

            contenido = archivo_principal.read_text(
                encoding="utf-8-sig"
            )

            bloque_original = (
                'archivos_csv = [\n'
                '    Path("pruebas_v2_06/archivo_X.csv"),\n'
                '    Path("pruebas_v2_06/archivo_W.csv")\n'
                ']'
            )

            bloque_prueba = (
                'archivos_csv = [\n'
                f'    Path("{archivo_A.as_posix()}"),\n'
                f'    Path("{archivo_B.as_posix()}")\n'
                ']'
            )

            self.assertIn(
                bloque_original,
                contenido,
                "No se encontró la configuración de entrada esperada."
            )

            contenido = contenido.replace(
                bloque_original,
                bloque_prueba,
                1
            )

            contenido = contenido.replace(
                '"salida/datos_limpios.csv"',
                f'"{(ruta_salida / "datos_limpios.csv").as_posix()}"'
            )

            contenido = contenido.replace(
                '"salida/informe_proceso.txt"',
                f'"{(ruta_salida / "informe_proceso.txt").as_posix()}"'
            )

            contenido = contenido.replace(
                '"salida/informe_proceso.xlsx"',
                f'"{(ruta_salida / "informe_proceso.xlsx").as_posix()}"'
            )

            script_temporal = carpeta / "inicio_prueba.py"

            script_temporal.write_text(
                contenido,
                encoding="utf-8"
            )

            entorno = os.environ.copy()
            entorno["PYTHONPATH"] = str(RAIZ_PROYECTO)
            entorno["PYTHONIOENCODING"] = "utf-8"

            resultado = subprocess.run(
                [sys.executable, str(script_temporal)],
                cwd=RAIZ_PROYECTO,
                capture_output=True,
                text=True,
                encoding="utf-8",
                errors="replace",
                env=entorno,
                timeout=60
            )

            salida_consola = (
                resultado.stdout + "\n" + resultado.stderr
            )

            self.assertEqual(
                resultado.returncode,
                0,
                "El programa terminó con error:\n"
                + salida_consola
            )

            archivo_limpio = (
                ruta_salida / "datos_limpios.csv"
            )

            informe_txt = (
                ruta_salida / "informe_proceso.txt"
            )

            informe_excel = (
                ruta_salida / "informe_proceso.xlsx"
            )

            self.assertTrue(
                archivo_limpio.is_file(),
                "No se generó el CSV limpio."
            )

            self.assertTrue(
                informe_txt.is_file(),
                "No se generó el informe TXT."
            )

            self.assertTrue(
                informe_excel.is_file(),
                "No se generó el informe Excel."
            )

            datos_limpios = pd.read_csv(archivo_limpio)

            self.assertEqual(
                len(datos_limpios),
                5,
                "El CSV limpio debe conservar los cinco registros."
            )

            self.assertIn(
                "Archivo_Origen",
                datos_limpios.columns,
                "No se conservó la procedencia de los registros."
            )

            self.assertEqual(
                datos_limpios["ID"].tolist(),
                [901, 902, 903, 904, 905]
            )

            self.assertEqual(
                datos_limpios["Nombre"].tolist(),
                [
                    "Producto ZZ",
                    "Producto ZZ",
                    "Producto ZZ",
                    "Producto ZZ",
                    "Producto YY"
                ]
            )

            contenido_txt = informe_txt.read_text(
                encoding="utf-8"
            )

            self.assertIn(
                "LISTO",
                contenido_txt,
                "El informe TXT no refleja el estado final LISTO."
            )

            self.assertIn(
                "Producto ZZ",
                contenido_txt,
                "El informe TXT no refleja la forma dominante."
            )

            self.assertIn(
                "espacios",
                contenido_txt.lower(),
                "El informe TXT no menciona las correcciones de espacios."
            )

            self.assertIn(
                "capitalización",
                contenido_txt.lower(),
                "El informe TXT no menciona las correcciones de capitalización."
            )

            libro = load_workbook(
                informe_excel,
                read_only=True,
                data_only=True
            )


            # Verificar los resultados reales de la validación
            # de capitalización dentro de su sección específica.
            self.assertIn(
                "VALIDACIÓN DE CORRECCIONES",
                contenido_txt,
                "Falta la sección de validación de correcciones."
            )

            seccion_validacion = contenido_txt.split(
                "VALIDACIÓN DE CORRECCIONES",
                1
            )[1]

            self.assertIn(
                "CONSISTENCIA DE TEXTO",
                seccion_validacion,
                "Falta la validación de consistencia de texto."
            )

            seccion_capitalizacion = (
                seccion_validacion
                .split("CONSISTENCIA DE TEXTO", 1)[1]
                .split("RESULTADO FINAL", 1)[0]
            )

            self.assertIn("Nombre:", seccion_capitalizacion)
            self.assertIn(
                "Grupo: producto zz",
                seccion_capitalizacion
            )
            self.assertIn(
                "Forma dominante: Producto ZZ",
                seccion_capitalizacion
            )
            self.assertIn("Antes: 1", seccion_capitalizacion)
            self.assertIn("Despues: 0", seccion_capitalizacion)
            self.assertIn("Corregidos: 1", seccion_capitalizacion)
            self.assertIn("Pendientes: 0", seccion_capitalizacion)
            self.assertIn("Estado: CORREGIDO", seccion_capitalizacion)

            try:
                hojas_esperadas = {
                    "Resumen",
                    "Calidad",
                    "Correcciones",
                    "Pendientes",
                    "Datos_Limpios"
                }

                self.assertTrue(
                    hojas_esperadas.issubset(
                        set(libro.sheetnames)
                    ),
                    "Faltan hojas esperadas en el informe Excel."
                )

                contenido_resumen = "\n".join(
                    str(valor)
                    for fila in libro["Resumen"].iter_rows(
                        values_only=True
                    )
                    for valor in fila
                    if valor is not None
                )

                self.assertIn(
                    "LISTO",
                    contenido_resumen,
                    "El resumen Excel no refleja el estado LISTO."
                )

                contenido_correcciones = "\n".join(
                    str(valor)
                    for fila in libro["Correcciones"].iter_rows(
                        values_only=True
                    )
                    for valor in fila
                    if valor is not None
                ).lower()

                self.assertIn(
                    "espacios",
                    contenido_correcciones,
                    "El Excel no documenta las correcciones de espacios."
                )

                self.assertIn(
                    "capitalizacion",
                    contenido_correcciones,
                    "El Excel no documenta las correcciones de capitalización."
                )

                hoja_datos = libro["Datos_Limpios"]

                encabezados = [
                    celda.value
                    for celda in hoja_datos[3]
                    if celda.value is not None
                ]

                self.assertEqual(
                    encabezados,
                    [
                        "ID",
                        "Nombre",
                        "Precio",
                        "Archivo_Origen"
                    ],
                    "Los encabezados de Datos_Limpios no son los esperados."
                )

                registros_excel = list(
                    hoja_datos.iter_rows(
                        min_row=4,
                        values_only=True
                    )
                )

                registros_excel = [
                    fila
                    for fila in registros_excel
                    if any(valor is not None for valor in fila)
                ]

                self.assertEqual(
                    len(registros_excel),
                    5,
                    "Datos_Limpios debe contener cinco registros."
                )

                identificadores_excel = [
                    fila[0]
                    for fila in registros_excel
                ]

                self.assertEqual(
                    identificadores_excel,
                    [901, 902, 903, 904, 905],
                    "Los identificadores de Datos_Limpios son incorrectos."
                )

                nombres_excel = [
                    fila[1]
                    for fila in registros_excel
                ]

                self.assertEqual(
                    nombres_excel,
                    [
                        "Producto ZZ",
                        "Producto ZZ",
                        "Producto ZZ",
                        "Producto ZZ",
                        "Producto YY"
                    ],
                    "Los nombres de Datos_Limpios no están normalizados."
                )

                procedencias_excel = [
                    fila[3]
                    for fila in registros_excel
                ]

                self.assertEqual(
                    procedencias_excel,
                    [
                        "prueba_A.csv",
                        "prueba_A.csv",
                        "prueba_A.csv",
                        "prueba_B.csv",
                        "prueba_B.csv"
                    ],
                    "La procedencia de los registros es incorrecta."
                )

            finally:
                libro.close()

if __name__ == "__main__":
    unittest.main(verbosity=2)
