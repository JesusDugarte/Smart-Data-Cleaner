
import os
import subprocess
import sys
import tempfile
import unittest
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook


RAIZ_PROYECTO = Path(__file__).resolve().parent.parent


class PruebaIntegracionTiposConvertibles(unittest.TestCase):

    def test_consolidacion_normal_de_principio_a_fin(self):

        archivos_entrada = [
            RAIZ_PROYECTO
            / "pruebas_v2_06"
            / "archivo_tipos_convertibles_A.csv",
            RAIZ_PROYECTO
            / "pruebas_v2_06"
            / "archivo_tipos_convertibles_B.csv",
        ]

        archivo_principal = RAIZ_PROYECTO / "inicio_v2_06.py"

        for archivo in archivos_entrada:
            self.assertTrue(
                archivo.is_file(),
                f"No existe el archivo de entrada: {archivo.name}",
            )

        self.assertTrue(
            archivo_principal.is_file(),
            "No existe inicio_v2_06.py.",
        )

        with tempfile.TemporaryDirectory(
            prefix="prueba_tipos_convertibles_"
        ) as carpeta_temporal:

            carpeta = Path(carpeta_temporal)
            ruta_salida = carpeta / "salida"
            ruta_salida.mkdir()

            contenido = archivo_principal.read_text(
                encoding="utf-8-sig"
            )

            bloque_original = (
                'archivos_csv = [\n'
                '    Path("pruebas_v2_06/archivo_X.csv"),\n'
                '    Path("pruebas_v2_06/archivo_W.csv")\n'
                ']'
            )

            bloque_prueba = (
                'archivos_csv = [\n'
                f'    Path("{archivos_entrada[0].as_posix()}"),\n'
                f'    Path("{archivos_entrada[1].as_posix()}")\n'
                ']'
            )

            self.assertIn(
                bloque_original,
                contenido,
                "No se encontró la configuración de entrada esperada.",
            )

            contenido = contenido.replace(
                bloque_original,
                bloque_prueba,
                1,
            )

            rutas_salida = {
                '"salida/datos_limpios.csv"':
                    ruta_salida / "datos_limpios.csv",
                '"salida/informe_proceso.txt"':
                    ruta_salida / "informe_proceso.txt",
                '"salida/informe_proceso.xlsx"':
                    ruta_salida / "informe_proceso.xlsx",
            }

            for ruta_original, ruta_temporal in rutas_salida.items():
                contenido = contenido.replace(
                    ruta_original,
                    f'"{ruta_temporal.as_posix()}"',
                )

            script_temporal = carpeta / "inicio_prueba.py"
            script_temporal.write_text(
                contenido,
                encoding="utf-8",
            )

            entorno = os.environ.copy()
            entorno["PYTHONPATH"] = str(RAIZ_PROYECTO)
            entorno["PYTHONIOENCODING"] = "utf-8"

            resultado = subprocess.run(
                [sys.executable, str(script_temporal)],
                cwd=RAIZ_PROYECTO,
                capture_output=True,
                text=True,
                encoding="utf-8",
                errors="replace",
                env=entorno,
                timeout=60,
            )

            salida_consola = (
                resultado.stdout + "\n" + resultado.stderr
            )

            self.assertEqual(
                resultado.returncode,
                0,
                "El programa terminó con error:\n" + salida_consola,
            )

            self.assertIn(
                "Estado de consolidación: REVISAR",
                salida_consola,
                "La consola no confirma el estado de consolidación REVISAR.",
            )

            self.assertIn(
                "Consolidación realizada: True",
                salida_consola,
                "La consola no confirma que se realizó la consolidación.",
            )

            self.assertIn(
                "COMPATIBILIDAD DE TIPOS",
                salida_consola,
                "La consola no muestra el diagnóstico de compatibilidad de tipos.",
            )

            self.assertIn(
                "Columna: Precio",
                salida_consola,
                "El diagnóstico no identifica la columna Precio.",
            )

            self.assertEqual(
                salida_consola.count("COMPATIBLE_POTENCIAL"),
                2,
                "Se esperaba COMPATIBLE_POTENCIAL para los dos archivos.",
            )

            self.assertNotIn(
                "Valores no convertibles",
                salida_consola,
                "La consola reportó valores no convertibles en este escenario.",
            )

            archivo_csv = ruta_salida / "datos_limpios.csv"
            informe_txt = ruta_salida / "informe_proceso.txt"
            informe_excel = ruta_salida / "informe_proceso.xlsx"

            for archivo in (archivo_csv, informe_txt, informe_excel):
                self.assertTrue(
                    archivo.is_file(),
                    f"No se generó el archivo esperado: {archivo.name}",
                )

            datos = pd.read_csv(archivo_csv)

            self.assertEqual(len(datos), 4)
            self.assertEqual(
                datos["ID"].tolist(),
                [1, 2, 3, 4],
            )

            self.assertIn("Archivo_Origen", datos.columns)

            self.assertEqual(
                datos["Archivo_Origen"].value_counts().to_dict(),
                {
                    "archivo_tipos_convertibles_A.csv": 2,
                    "archivo_tipos_convertibles_B.csv": 2,
                },
            )

            contenido_txt = informe_txt.read_text(encoding="utf-8")

            self.assertIn(
                "Estado final: REVISAR",
                contenido_txt,
                "El informe TXT no confirma el estado final REVISAR.",
            )

            libro = load_workbook(
                informe_excel,
                read_only=True,
                data_only=True,
            )

            try:
                hojas_esperadas = {
                    "Resumen",
                    "Calidad",
                    "Correcciones",
                    "Pendientes",
                    "Datos_Limpios",
                }

                self.assertTrue(
                    hojas_esperadas.issubset(set(libro.sheetnames)),
                    "Faltan hojas esperadas en el informe Excel.",
                )

                contenido_resumen = "\n".join(
                    str(valor)
                    for fila in libro["Resumen"].iter_rows(
                        values_only=True
                    )
                    for valor in fila
                    if valor is not None
                )

                self.assertIn(
                    "REVISAR",
                    contenido_resumen,
                    "El resumen Excel no confirma el estado REVISAR.",
                )

            finally:
                libro.close()


if __name__ == "__main__":
    unittest.main(verbosity=2)
