
import unittest
from unittest.mock import Mock, patch

from openpyxl import Workbook

from exportacion.excel import (
    construir_resumen,
    guardar_libro_informe
)

class PruebaExportacionExcel(unittest.TestCase):

    def crear_resultado_proceso(self, resultados_procedencia):
        return {
            "resultado": {
                "resumen": {
                    "archivos_analizados": 2,
                    "registros_procesados": 4
                },
                "estado_final": "LISTO"
            },
            "ejecucion": {
                "consolidacion_realizada": True
            },
            "validacion": {
                "registros": {
                    "valida": True
                },
                "procedencia": {
                    "resultados": resultados_procedencia
                },
                "archivo_exportado": {
                    "estado": "Correcto"
                }
            }
        }

    def test_error_de_permisos_se_propaga(self):
        libro = Mock()

        with patch.object(
            libro,
            "save",
            side_effect=PermissionError("Acceso denegado")
        ):
            with self.assertRaisesRegex(
                PermissionError,
                "Acceso denegado"
            ):
                guardar_libro_informe(
                    libro,
                    "informe_prueba.xlsx"
                )

            libro.save.assert_called_once_with(
                "informe_prueba.xlsx"
            )

    def test_procedencia_sin_resultados_se_informa_correctamente(self):

        libro = Workbook()
        hoja = libro.active

        resultado_proceso = self.crear_resultado_proceso({})

        construir_resumen(hoja, resultado_proceso)

        self.assertEqual(
            hoja["B14"].value,
            "Sin resultados"
        )

    def test_procedencia_no_realizada(self):
        libro = Workbook()
        hoja = libro.active

        resultado_proceso = self.crear_resultado_proceso(None)

        construir_resumen(hoja, resultado_proceso)

        self.assertEqual(
            hoja["B14"].value,
            "No realizada"
        )

    def test_procedencia_correcta(self):
        libro = Workbook()
        hoja = libro.active

        resultados = {
            "archivo_A.csv": {
                "valida": True
            },
            "archivo_B.csv": {
                "valida": True
            }
        }

        resultado_proceso = self.crear_resultado_proceso(resultados)

        construir_resumen(hoja, resultado_proceso)

        self.assertEqual(
            hoja["B14"].value,
            "Correcto"
        )

    def test_procedencia_incorrecta(self):
        libro = Workbook()
        hoja = libro.active

        resultados = {
            "archivo_A.csv": {
                "valida": True
            },
            "archivo_B.csv": {
                "valida": False
            }
        }

        resultado_proceso = self.crear_resultado_proceso(resultados)

        construir_resumen(hoja, resultado_proceso)

        self.assertEqual(
            hoja["B14"].value,
            "Incorrecto"
        )

if __name__ == "__main__":
    unittest.main(verbosity=2)
