
import os
import subprocess
import sys
import tempfile
import unittest
from pathlib import Path

from openpyxl import load_workbook


RAIZ_PROYECTO = Path(__file__).resolve().parent.parent


class PruebaIntegracionEstructuraParcial(unittest.TestCase):

    def test_estructura_parcial_consolida_y_genera_informes(self):
        archivo_a = (
            RAIZ_PROYECTO
            / "pruebas_v2_06"
            / "prueba_estructura_parcial_A.csv"
        )
        archivo_b = (
            RAIZ_PROYECTO
            / "pruebas_v2_06"
            / "prueba_estructura_parcial_B.csv"
        )
        archivo_principal = RAIZ_PROYECTO / "inicio_v2_06.py"

        for archivo in (archivo_a, archivo_b, archivo_principal):
            self.assertTrue(
                archivo.is_file(),
                f"No existe el archivo requerido: {archivo}"
            )

        with tempfile.TemporaryDirectory(
            prefix="prueba_estructura_parcial_"
        ) as carpeta_temporal:

            carpeta = Path(carpeta_temporal)
            ruta_salida = carpeta.as_posix()

            contenido = archivo_principal.read_text(
                encoding="utf-8-sig"
            )

            bloque_original = (
                'archivos_csv = [\n'
                '    Path("pruebas_v2_06/archivo_X.csv"),\n'
                '    Path("pruebas_v2_06/archivo_W.csv")\n'
                ']'
            )

            bloque_prueba = (
                'archivos_csv = [\n'
                f'    Path("{archivo_a.as_posix()}"),\n'
                f'    Path("{archivo_b.as_posix()}")\n'
                ']'
            )

            self.assertIn(
                bloque_original,
                contenido,
                "No se encontró la configuración de entrada esperada."
            )

            contenido = contenido.replace(
                bloque_original,
                bloque_prueba,
                1
            )

            contenido = contenido.replace(
                '"salida/datos_limpios.csv"',
                f'"{ruta_salida}/datos_limpios.csv"'
            )
            contenido = contenido.replace(
                'Path("salida/datos_limpios.csv")',
                f'Path("{ruta_salida}/datos_limpios.csv")'
            )
            contenido = contenido.replace(
                '"salida/informe_proceso.txt"',
                f'"{ruta_salida}/informe_proceso.txt"'
            )
            contenido = contenido.replace(
                '"salida/informe_proceso.xlsx"',
                f'"{ruta_salida}/informe_proceso.xlsx"'
            )

            script_temporal = carpeta / "inicio_prueba.py"
            script_temporal.write_text(
                contenido,
                encoding="utf-8"
            )

            ejecucion = subprocess.run(
                [sys.executable, str(script_temporal)],
                cwd=RAIZ_PROYECTO,
                capture_output=True,
                text=True,
                encoding="utf-8",
                errors="replace",
                env={
                    **os.environ,
                    "PYTHONPATH": (
                        str(RAIZ_PROYECTO)
                        + os.pathsep
                        + os.environ.get("PYTHONPATH", "")
                    ),
                    "PYTHONIOENCODING": "utf-8",
                },
            )

            salida_consola = (
                ejecucion.stdout + ejecucion.stderr
            )

            self.assertEqual(
                ejecucion.returncode,
                0,
                f"El programa terminó con error:\n{salida_consola}"
            )

            self.assertIn(
                "Estado final: REVISAR",
                salida_consola
            )
            self.assertIn(
                "Consolidación realizada: True",
                salida_consola
            )
            self.assertIn(
                "Esperados: 4",
                salida_consola
            )
            self.assertIn(
                "Obtenidos: 4",
                salida_consola
            )
            self.assertIn(
                "Descuento",
                salida_consola
            )

            archivo_csv = carpeta / "datos_limpios.csv"
            archivo_txt = carpeta / "informe_proceso.txt"
            archivo_excel = carpeta / "informe_proceso.xlsx"

            self.assertTrue(archivo_csv.is_file())
            self.assertTrue(archivo_txt.is_file())
            self.assertTrue(archivo_excel.is_file())

            import pandas as pd

            datos = pd.read_csv(archivo_csv)

            self.assertEqual(len(datos), 4)
            self.assertEqual(len(datos.columns), 5)
            self.assertEqual(
                datos["Descuento"].isna().sum(),
                2
            )

            texto_informe = archivo_txt.read_text(
                encoding="utf-8"
            )

            self.assertIn("REVISAR", texto_informe)
            self.assertIn("Descuento", texto_informe)
            self.assertIn("revisar_faltantes", texto_informe)

            libro = load_workbook(
                archivo_excel,
                read_only=True,
                data_only=True
            )

            try:
                self.assertIn("Resumen", libro.sheetnames)
                self.assertIn("Calidad", libro.sheetnames)
                self.assertIn("Pendientes", libro.sheetnames)

                valores_excel = [
                    str(celda.value)
                    for hoja in libro.worksheets
                    for fila in hoja.iter_rows()
                    for celda in fila
                    if celda.value is not None
                ]

                self.assertIn("REVISAR", valores_excel)
                self.assertIn("Descuento", valores_excel)
                self.assertIn("revisar_faltantes", valores_excel)
            finally:
                libro.close()


if __name__ == "__main__":
    unittest.main(verbosity=2)
