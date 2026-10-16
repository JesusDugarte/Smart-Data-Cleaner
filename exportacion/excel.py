from openpyxl import Workbook
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.cell.cell import MergedCell
import pandas as pd

def crear_libro_informe():

    libro = Workbook()

    libro.active.title = "Resumen"

    libro.create_sheet("Calidad")
    libro.create_sheet("Correcciones")
    libro.create_sheet("Pendientes")
    libro.create_sheet("Datos_Limpios")

    return libro

def construir_resumen(hoja, resultado_proceso):

    resumen = resultado_proceso["resultado"]["resumen"]
    estado_final = resultado_proceso["resultado"]["estado_final"]
    ejecucion = resultado_proceso["ejecucion"]
    validacion = resultado_proceso["validacion"]

    hoja["A1"] = "SMART DATA CLEANER"
    hoja["A2"] = "INFORME DE CALIDAD DE DATOS"

    hoja["A4"] = "RESUMEN DEL PROCESO"

    hoja["A5"] = "Archivos analizados"
    hoja["B5"] = resumen["archivos_analizados"]

    hoja["A6"] = "Registros procesados"
    hoja["B6"] = resumen["registros_procesados"]

    hoja["A7"] = "Estado final"
    hoja["B7"] = estado_final

    hoja["A9"] = "EJECUCIÓN"


    hoja["A10"] = "Consolidación realizada"

    consolidacion_realizada = ejecucion["consolidacion_realizada"]

    if consolidacion_realizada is True:
        consolidacion_realizada = "Sí"
    elif consolidacion_realizada is False:
        consolidacion_realizada = "No"

    hoja["B10"] = consolidacion_realizada


    hoja["A12"] = "VALIDACIÓN"

    resultado_registros = validacion["registros"]["valida"]

    if resultado_registros is None:
        resultado_registros = "No realizada"
    elif resultado_registros is True:
        resultado_registros = "Correcto"
    else:
        resultado_registros = "Incorrecto"

    hoja["A13"] = "Registros"
    hoja["B13"] = resultado_registros

    resultados_procedencia = (
        validacion["procedencia"]["resultados"]
    )



    if resultados_procedencia is None:
        validacion_procedencia = "No realizada"

    elif not resultados_procedencia:
        validacion_procedencia = "Sin resultados"

    else:
        validacion_procedencia = all(
            resultado["valida"]
            for resultado in resultados_procedencia.values()
        )

        if validacion_procedencia is True:
            validacion_procedencia = "Correcto"
        else:
            validacion_procedencia = "Incorrecto"



    hoja["A14"] = "Procedencia"
    hoja["B14"] = validacion_procedencia

    archivo_exportado = validacion["archivo_exportado"]

    hoja["A15"] = "Archivo exportado"
    hoja["B15"] = archivo_exportado["estado"]

def formatear_resumen(hoja):

    # Formato visual de la identidad del informe
    # Combinar las celdas de los títulos principales
    hoja.merge_cells("A1:B1")
    hoja.merge_cells("A2:B2")
    hoja["A1"].font = Font(
        bold=True,
        size=16,
        color="FFFFFF"
    )

    hoja["A1"].fill = PatternFill(
        fill_type="solid",
        fgColor="17365D"
    )

    hoja["A2"].font = Font(
        bold=True,
        size=12,
        color="FFFFFF"
    )

    hoja["A2"].fill = PatternFill(
        fill_type="solid",
        fgColor="17365D"
    )


    # Definir un borde fino y discreto
    borde_fino = Side(
        style="thin",
        color="B4C6E7"
    )

    borde_seccion = Border(
        bottom=borde_fino
    )


    # Bordes discretos para el resumen principal
    borde_detalle = Side(
        style="thin",
        color="D9E2F3"
    )


    # Formato visual de las secciones
    for fila in (4, 9, 12):
        for columna in ("A", "B"):
            celda = hoja[f"{columna}{fila}"]

            celda.fill = PatternFill(
                fill_type="solid",
                fgColor="D9EAF7"
            )
            celda.border = borde_seccion

        hoja[f"A{fila}"].font = Font(
            bold=True,
            color="17365D"
        )

    # Delimitar el bloque principal del resumen
    for fila in range(5, 8):
        hoja[f"A{fila}"].border = Border(
            left=borde_detalle,
            bottom=borde_detalle if fila == 7 else Side(style=None)
        )

        hoja[f"B{fila}"].border = Border(
            right=borde_detalle,
            bottom=borde_detalle if fila == 7 else Side(style=None)
        )

    hoja["A5"].border = Border(
        left=borde_detalle,
        top=borde_detalle
    )

    hoja["B5"].border = Border(
        right=borde_detalle,
        top=borde_detalle
    )

    hoja["B5"].alignment = Alignment(
        horizontal="center"
    )

    hoja["B6"].alignment = Alignment(
        horizontal="center"
    )

    hoja["B7"].alignment = Alignment(
        horizontal="center"
    )

    hoja["B10"].alignment = Alignment(
        horizontal="center"
    )

    hoja["B13"].alignment = Alignment(
        horizontal="center"
    )

    hoja["B14"].alignment = Alignment(
        horizontal="center"
    )

    hoja["B15"].alignment = Alignment(
        horizontal="center"
    )

    estado = hoja["B7"].value

    if estado == "LISTO":

        hoja["B7"].fill = PatternFill(
            fill_type="solid",
            fgColor="C6EFCE"
        )

    elif estado == "REVISAR":

        hoja["B7"].fill = PatternFill(
            fill_type="solid",
            fgColor="FFEB9C"
        )

    elif estado == "BLOQUEADO":

        hoja["B7"].fill = PatternFill(
            fill_type="solid",
            fgColor="FFC7CE"
        )

    # Formato visual de las validaciones
    for fila in (13, 14):
        celda = hoja[f"B{fila}"]
        resultado = celda.value

        if resultado == "Correcto":
            celda.fill = PatternFill(
                fill_type="solid",
                fgColor="C6EFCE"
            )
        elif resultado == "Incorrecto":
            celda.fill = PatternFill(
                fill_type="solid",
                fgColor="FFC7CE"
            )

def construir_calidad(hoja, resultado_calidad):

    hoja["A1"] = "CALIDAD DE DATOS"

    if not resultado_calidad["ejecucion"]["calidad"]["realizada"]:

        hoja["A3"] = "Análisis de calidad no realizado."

        return

    construir_calidad_faltantes(
        hoja,
        resultado_calidad["diagnostico"],
        resultado_calidad["decision"]["calidad"]
    )

    construir_calidad_duplicados(
        hoja,
        resultado_calidad["diagnostico"]
    )

    construir_calidad_espacios(
        hoja,
        resultado_calidad["diagnostico"]
    )

    construir_calidad_consistencia(
        hoja,
        resultado_calidad["diagnostico"],
        resultado_calidad["decision"]["calidad"]
    )

def construir_calidad_faltantes(
    hoja,
    diagnostico_calidad,
    decision_calidad
):

    faltantes = diagnostico_calidad["faltantes"]

    clasificaciones = {}

    for hallazgo in decision_calidad["hallazgos"]:

        if hallazgo["categoria"] == "faltantes":

            clasificaciones[hallazgo["columna"]] = hallazgo["problema"]

    hoja["A1"] = "CALIDAD DE DATOS"

    hoja["A3"] = "FALTANTES"

    hoja["A4"] = "Columna"
    hoja["B4"] = "Total"
    hoja["C4"] = "Presentes"
    hoja["D4"] = "Faltantes"
    hoja["E4"] = "% faltante"
    hoja["F4"] = "Clasificación"

    fila = 5

    for columna, detalle in faltantes.items():

        hoja.cell(fila, 1).value = columna
        hoja.cell(fila, 2).value = detalle["total"]
        hoja.cell(fila, 3).value = detalle["presentes"]
        hoja.cell(fila, 4).value = detalle["faltantes"]
        hoja.cell(fila, 5).value = detalle["porcentaje_faltante"]
        if detalle["faltantes"] == 0:
            clasificacion = "SIN_FALTANTES"
        else:
            clasificacion = clasificaciones.get(
                columna,
                "SIN_CLASIFICAR"
            )

        etiquetas_clasificacion = {
            "SIN_FALTANTES": "Sin faltantes",
            "SIN_CLASIFICAR": "Sin clasificar"
        }

        clasificacion_visible = etiquetas_clasificacion.get(
            clasificacion,
            clasificacion.replace("_", " ").capitalize()
        )

        hoja.cell(fila, 6).value = clasificacion_visible


        fila += 1

def construir_calidad_duplicados(hoja, diagnostico_calidad):

    duplicados = diagnostico_calidad["duplicados"]

    hoja["A10"] = "DUPLICADOS"

    hoja["A11"] = "Total"
    hoja["B11"] = "Filas duplicadas"
    hoja["C11"] = "Filas únicas"
    hoja["D11"] = "Porcentaje"

    hoja["A12"] = duplicados["total"]
    hoja["B12"] = duplicados["filas_duplicadas"]
    hoja["C12"] = duplicados["filas_unicas"]
    hoja["D12"] = duplicados["porcentaje_duplicados"]

def construir_calidad_espacios(hoja, diagnostico_calidad):

    espacios = diagnostico_calidad["espacios"]

    hoja["A15"] = "ESPACIOS"

    hoja["A16"] = "Columna"
    hoja["B16"] = "Total"
    hoja["C16"] = "Valores afectados"
    hoja["D16"] = "Iniciales"
    hoja["E16"] = "Finales"
    hoja["F16"] = "Internos"

    fila = 17

    for columna, detalle in espacios.items():

        hoja.cell(fila, 1).value = columna
        hoja.cell(fila, 2).value = detalle["total"]
        hoja.cell(fila, 3).value = detalle["valores_con_espacios"]
        hoja.cell(fila, 4).value = detalle["espacios_iniciales"]
        hoja.cell(fila, 5).value = detalle["espacios_finales"]
        hoja.cell(fila, 6).value = detalle["espacios_internos"]

        fila += 1

def construir_calidad_consistencia(
    hoja,
    diagnostico_calidad,
    decision_calidad
):

    consistencia = diagnostico_calidad["consistencia_texto"]

    hoja["A20"] = "CONSISTENCIA DE TEXTO"

    hoja["A21"] = "Columna"
    hoja["B21"] = "Grupo"
    hoja["C21"] = "Variaciones"

    fila = 22

    for columna, detalle in consistencia.items():

        variaciones = detalle["detalle_variaciones"]

        for grupo in variaciones:

            hoja.cell(fila, 1).value = columna
            hoja.cell(fila, 2).value = grupo
            hoja.cell(fila, 3).value = (
                detalle["variaciones_capitalizacion"]
            )

            fila += 1

    hoja["A25"] = "FORMAS OBSERVADAS"

    hoja["A26"] = "Columna"
    hoja["B26"] = "Grupo"
    hoja["C26"] = "Forma observada"
    hoja["D26"] = "Frecuencia"

    fila = 27

    for columna, detalle in consistencia.items():

        frecuencias = detalle["frecuencias_variaciones"]

        for grupo, variantes in frecuencias.items():

            for forma, frecuencia in variantes.items():

                hoja.cell(fila, 1).value = columna
                hoja.cell(fila, 2).value = grupo
                hoja.cell(fila, 3).value = forma
                hoja.cell(fila, 4).value = frecuencia

                fila += 1

    hoja["A31"] = "DECISIÓN DE CAPITALIZACIÓN"

    hoja["A32"] = "Grupo"
    hoja["B32"] = "Forma dominante"
    hoja["C32"] = "% dominante"
    hoja["D32"] = "Corrección posible"

    fila = 33

    for hallazgo in decision_calidad["hallazgos"]:

        if hallazgo["categoria"] != "consistencia_texto":
            continue

        decisiones = hallazgo["detalle"]["decisiones"]

        for grupo, decision in decisiones.items():

            hoja.cell(fila, 1).value = grupo
            hoja.cell(fila, 2).value = decision["forma_dominante"]
            hoja.cell(fila, 3).value = decision["porcentaje"]
            hoja.cell(fila, 3).number_format = '0.00"%"'
            hoja.cell(fila, 4).value = (
                "Sí"
                if decision["corregible"]
                else "No"
            )

            fila += 1

def formatear_calidad(hoja):


    # Título principal del informe de calidad
    hoja["A1"].font = Font(
        bold=True,
        size=14,
        color="FFFFFF"
    )

    hoja["A1"].fill = PatternFill(
        fill_type="solid",
        fgColor="17365D"
    )

    # Bandas visuales para las secciones
    for fila in (3, 10, 15, 20, 25, 31):
        for columna in range(1, 7):
            celda = hoja.cell(fila, columna)

            if celda.value is not None or columna == 1:
                celda.fill = PatternFill(
                    fill_type="solid",
                    fgColor="D9EAF7"
                )

        hoja[f"A{fila}"].font = Font(
            bold=True,
            color="17365D"
        )

# Encabezados de las tablas
    filas_encabezado = [
        4,
        11,
        16,
        21,
        26,
        32
    ]

    for fila in filas_encabezado:

        for celda in hoja[fila]:

            if celda.value is not None:

                celda.font = Font(
                    bold=True
                )

                celda.alignment = Alignment(
                    horizontal="center"
                )

    # Formato de porcentajes
    hoja["E5"].number_format = '0.00"%"'
    hoja["E6"].number_format = '0.00"%"'
    hoja["E7"].number_format = '0.00"%"'

    hoja["D12"].number_format = '0.00"%"'

    # Porcentajes de decisión de capitalización
    for fila in range(33, hoja.max_row + 1):

        if hoja.cell(fila, 3).value is not None:

            hoja.cell(
                fila,
                3
            ).number_format = '0.00"%"'

def construir_correcciones(hoja, resultado_proceso):

    correcciones = resultado_proceso["resultado"]["correcciones"]

    hoja["A1"] = "CORRECCIONES REALIZADAS"

    hoja["A3"] = "Categoría"
    hoja["B3"] = "Columna"
    hoja["C3"] = "Valores corregidos"

    fila = 4

    for categoria, detalles in correcciones.items():

        for columna, detalle in detalles.items():

            hoja.cell(fila, 1).value = categoria
            hoja.cell(fila, 2).value = columna
            hoja.cell(fila, 3).value = (
                detalle["valores_corregidos"]
            )

            fila += 1

def formatear_correcciones(hoja):

    # Combinar las celdas del título principal
    hoja.merge_cells("A1:C1")
    # Título principal de la hoja
    hoja["A1"].font = Font(
        bold=True,
        size=14,
        color="FFFFFF"
    )

    hoja["A1"].fill = PatternFill(
        fill_type="solid",
        fgColor="17365D"
    )



    # Formato de los encabezados de la tabla
    for celda in hoja[3]:

        if celda.value is not None:

            celda.font = Font(
                bold=True,
                color="17365D"
            )

            celda.fill = PatternFill(
                fill_type="solid",
                fgColor="D9EAF7"
            )

            celda.alignment = Alignment(
                horizontal="center",
                vertical="center"
            )


    for fila in range(4, hoja.max_row + 1):

        hoja.cell(
            fila,
            3
        ).alignment = Alignment(
            horizontal="center"
        )

def construir_pendientes(hoja, resultado_proceso):

    pendientes = resultado_proceso["resultado"]["pendientes"]

    hoja["A1"] = "PENDIENTES"

    hoja["A3"] = "Categoría"
    hoja["B3"] = "Columna"
    hoja["C3"] = "Acción"

    fila = 4

    for pendiente in pendientes:

        hoja.cell(fila, 1).value = pendiente["categoria"]
        hoja.cell(fila, 2).value = pendiente["columna"]
        hoja.cell(fila, 3).value = pendiente["accion"]

        fila += 1

def formatear_pendientes(hoja):

    # Título principal
    hoja.merge_cells("A1:C1")

    hoja["A1"].font = Font(
        bold=True,
        size=14,
        color="FFFFFF"
    )

    hoja["A1"].fill = PatternFill(
        fill_type="solid",
        fgColor="17365D"
    )



    # Formato de los encabezados
    for celda in hoja[3]:

        if celda.value is not None:

            celda.font = Font(
                bold=True,
                color="17365D"
            )

            celda.fill = PatternFill(
                fill_type="solid",
                fgColor="D9EAF7"
            )

            celda.alignment = Alignment(
                horizontal="center",
                vertical="center"
            )


    for fila in range(4, hoja.max_row + 1):

        hoja.cell(fila, 3).alignment = Alignment(
            horizontal="left"
        )

def construir_datos_limpios(hoja, resultado_proceso):

    datos = resultado_proceso["resultado"]["datos"]

    hoja["A1"] = "DATOS LIMPIOS"

    for columna_numero, columna in enumerate(datos.columns, start=1):
        hoja.cell(3, columna_numero).value = columna

    for fila_numero, fila_datos in enumerate(
        datos.itertuples(index=False, name=None),
        start=4
    ):
        for columna_numero, valor in enumerate(
            fila_datos,
            start=1
        ):
            valor_excel = None if pd.isna(valor) else valor

            hoja.cell(
                fila_numero,
                columna_numero
            ).value = valor_excel


def formatear_datos_limpios(hoja):

    # Combinar las celdas del título principal
    hoja.merge_cells(
        start_row=1,
        start_column=1,
        end_row=1,
        end_column=max(hoja.max_column, 3)
    )

    hoja["A1"].font = Font(
        bold=True,
        size=14,
        color="FFFFFF"
    )

    hoja["A1"].fill = PatternFill(
        fill_type="solid",
        fgColor="17365D"
    )


    for celda in hoja[3]:
        celda.font = Font(
            bold=True
        )

    for celda in hoja[3]:
        celda.alignment = Alignment(
            horizontal="center"
        )

    hoja.freeze_panes = "A4"

    if hoja.max_row >= 3 and hoja.max_column >= 1 and hoja["A3"].value is not None:

        tabla = Table(
            displayName="TablaDatosLimpios",
            ref=f"A3:{hoja.cell(hoja.max_row, hoja.max_column).coordinate}"
        )

        estilo_tabla = TableStyleInfo(
            name="TableStyleMedium2",
            showFirstColumn=False,
            showLastColumn=False,
            showRowStripes=True,
            showColumnStripes=False
        )

        tabla.tableStyleInfo = estilo_tabla

        hoja.add_table(tabla)
    
def ajustar_ancho_columnas(hoja):

    for columna in hoja.columns:

        ancho_maximo = 0
        letra_columna = None

        for celda in columna:

            if isinstance(celda, MergedCell):
                continue

            if letra_columna is None:
                letra_columna = celda.column_letter

            if celda.value is None:
                continue

            longitud = len(str(celda.value))

            if longitud > ancho_maximo:
                ancho_maximo = longitud

        if letra_columna is None:
            continue


        ancho_final = max(
            ancho_maximo + 2,
            12
        )

        ancho_final = min(
            ancho_final,
            30
        )

        hoja.column_dimensions[letra_columna].width = ancho_final


def guardar_libro_informe(libro, ruta_salida):

    libro.save(ruta_salida)


if __name__ == "__main__":

    libro = crear_libro_informe()

    guardar_libro_informe(
        libro,
        "salida/informe_proceso.xlsx"
    )

    print(libro.sheetnames)